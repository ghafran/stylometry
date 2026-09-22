"""Build the editable research manuscript from the reviewed Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(1.0)
section.left_margin = section.right_margin = Inches(0.9)
section.footer_distance = Inches(0.4)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
for name, size in [('Title', 19), ('Heading 1', 14), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != 'Title'
    style.paragraph_format.space_before = Pt(13 if name != 'Title' else 0)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_with_next = True
doc.styles['Title'].paragraph_format.space_after = Pt(16)
caption = doc.styles['Caption']
caption.font.name = 'Times New Roman'
caption.font.size = Pt(10)
caption.font.italic = False
caption.font.bold = True
caption.font.color.rgb = RGBColor(0, 0, 0)
caption.paragraph_format.keep_with_next = True
caption.paragraph_format.space_before = Pt(6)
caption.paragraph_format.space_after = Pt(5)

# Remove inherited decorative rules and theme fonts from the base template.
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for fonts in style.element.iter(qn('w:rFonts')):
        for name in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn('w:' + name), None)

def add_text(paragraph, text):
    pos = 0
    for m in re.finditer(r'\[([^\]]+)\]\((https?://[^\s]+)\)', text):
        paragraph.add_run(text[pos:m.start()])
        rel = paragraph.part.relate_to(m.group(2), RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), rel)
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79'); props.append(color)
        under = OxmlElement('w:u'); under.set(qn('w:val'), 'single'); props.append(under)
        run.append(props)
        t = OxmlElement('w:t'); t.text = m.group(1); run.append(t)
        link.append(run); paragraph._p.append(link)
        pos = m.end()
    paragraph.add_run(text[pos:])

def add_table(lines, number):
    rows = [[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    rows = [rows[0]] + rows[2:]
    widths = {1: [1.10, 1.02, 1.22, 1.03, 1.15, 1.18], 2: [1.20, .63, 1.46, .81, 1.55, .65], 3: [2.10, .93, .75, .96, .96, 1.0]}[number]
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths): col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
    table._tbl.tblPr.append(borders)
    for ridx, contents in enumerate(rows):
        cells = table.add_row().cells
        rowprops = table.rows[-1]._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); rowprops.append(cant)
        if ridx == 0:
            repeat = OxmlElement('w:tblHeader'); rowprops.append(repeat)
        for cidx, (cell, text, width) in enumerate(zip(cells, contents, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'bottom', 'left', 'right']:
                el = OxmlElement('w:' + edge); el.set(qn('w:w'), '75'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            props.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EDF3' if ridx == 0 else 'FFFFFF'); props.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text); run.font.size = Pt(9.5); run.bold = ridx == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

lines = (ROOT / 'stylometry_research_article.md').read_text().splitlines()
i = 0
table_number = 0
in_refs = False
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line.startswith('|'):
        block = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i]); i += 1
        table_number += 1; add_table(block, table_number); continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], style='Title')
    elif line.startswith('## '):
        heading = line[3:]
        in_refs = heading == 'References'
        doc.add_paragraph(heading, style='Heading 1')
    elif line.startswith('### '):
        heading = re.sub(r'^\d+\.\d+ ', '', line[4:])
        doc.add_paragraph(heading, style='Heading 2')
    elif line.startswith('Table '):
        doc.add_paragraph(line, style='Caption')
    else:
        p = doc.add_paragraph()
        add_text(p, line)
        if in_refs:
            p.paragraph_format.left_indent = Inches(.22)
            p.paragraph_format.first_line_indent = Inches(-.22)
            p.paragraph_format.keep_together = True
        if line.startswith('Keywords:'):
            for run in p.runs: run.font.size = Pt(10)
    i += 1

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'Style Discovery and Authorship Attribution in a Multilingual Corpus of Scriptural and Literary Texts'
doc.core_properties.subject = 'Research article based on the saved Stylometry 0.2.0 analysis'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.save(ROOT / 'stylometry_research_article.docx')
print(ROOT / 'stylometry_research_article.docx')
